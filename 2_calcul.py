"""
2_calcul.py — À partir du cache d'extraction, calcule une estimation de l'empreinte
carbone du transport des œuvres selon la norme ISO 14083, et produit un classeur Excel.

Formule ISO 14083 :  émissions = distance (km) × masse (t) × facteur d'émission (kgCO2e/t.km)

Lancement : uv run 2_calcul.py  (après avoir lancé 1_extraction.py)
"""

import json
import math
from pathlib import Path

from openpyxl import Workbook

# =============================================================================

FICHIER_DONNEES = Path("donnees") / "mouvements_bruts.json"
FICHIER_SORTIE = Path("sorties") / "resultats.xlsx"

# Siège du Frac Picardie (Amiens).
SIEGE_LAT, SIEGE_LON = 49.886, 2.31153

# --- Distances ---------------------------------------------------------------
# Aérien (EN 16258, §distance) : orthodromie + 95 km par vol.
SUPPLEMENT_AERIEN_KM = 95

# Routier : facteurs de détour nationaux (Ballou et al. 2002, Table 1).
# La distance réelle est approchée par : orthodromie × facteur du pays de destination.
FACTEUR_DETOUR = {
    "France": 1.65,
    "Italie": 1.18,
    "Allemagne": 1.32,
    "Espagne": 1.58,
    "Royaume-Uni": 1.40,
    "Pologne": 1.21,
    "Hongrie": 1.35,
}
FACTEUR_DETOUR_EUROPE = 1.46  # moyenne « Europe » de Ballou (n=199), pour les pays absents

# --- Masse (forfaits définis).
MASSE_OEUVRE_NUE_KG = {
    "Dessin": 4,
    "Estampe": 4,
    "Photographie": 4,
    "Peinture": 20,
    "Sculpture": 60,
    "Oeuvre en 3 dimensions": 60,
    "Installation": 80,
}
MASSE_OEUVRE_DEFAUT_KG = 15   # si le domaine n'est pas dans la liste

# Conditionnement : +30 % à l'oeuvre nue — convention reprise du calculateur GCC

MAJORATION_CONDITIONNEMENT = 1.30

# --- Mode de transport (hypothèses) ---
# Règle d'imputation : Europe = routier / hors Europe = aérien.

PAYS_EUROPE = {
    "France", "Allemagne", "Belgique", "Pays-Bas", "Espagne", "Italie", "Portugal",
    "Suisse", "Royaume-Uni", "Luxembourg", "Autriche", "Danemark", "Irlande",
    "Pologne", "Suède", "Norvège", "Finlande", "Grèce", "République tchèque", "Hongrie",
}

# --- Facteurs d'émission (kgCO2e par tonne-kilomètre) — ADEME Base Carbone® v23.11 ---
# Source : Base_Carbone_V23.11
# Routier : « Articulé, 34 à 40 tonnes, Diesel routier » = 0,0875 kgCO2e/t.km
# Aérien  : « Avion cargo, plus de 100 tonnes, >5000 kms, 2023, AVEC traînées » = 1,01 kgCO2e/t.km

FE_KG_PAR_TKM = {
    "Routier": 0.0875,
    "Aérien": 1.01,
}

# =============================================================================
# FONCTIONS
# =============================================================================

def charger_mouvements() -> list[dict]:
    """Lit le cache JSON et renvoie la liste des dictionnaires de mouvement."""
    donnees = json.loads(FICHIER_DONNEES.read_text(encoding="utf-8"))
    mouvements = []
    for resultat in donnees.get("results", []):
        bloc = resultat.get("_source", {}).get("ua", {}).get("movements", {})
        if bloc:
            mouvements.append(bloc)
    return mouvements


def haversine_km(lat1, lon1, lat2, lon2) -> float:
    """Distance orthodromique (km) entre deux points (formule de Haversine)."""
    rayon_terre = 6371.0  # km
    d_lat = math.radians(lat2 - lat1)
    d_lon = math.radians(lon2 - lon1)
    a = (
        math.sin(d_lat / 2) ** 2
        + math.cos(math.radians(lat1)) * math.cos(math.radians(lat2)) * math.sin(d_lon / 2) ** 2
    )
    return 2 * rayon_terre * math.asin(math.sqrt(a))


def imputer_masse_kg(domaines: str) -> float:
    """Masse imputée d'UNE œuvre (kg) = forfait « œuvre nue » × 1,30 conditionnement (GCC).

    Le domaine retenu est le premier de la liste `concerned_domains` (domaine dominant
    du mouvement) : la donnée Navigart ne fournit pas le domaine œuvre par œuvre.
    """
    premier = (domaines or "").split(",")[0].strip()
    nue = MASSE_OEUVRE_NUE_KG.get(premier, MASSE_OEUVRE_DEFAUT_KG)
    return round(nue * MAJORATION_CONDITIONNEMENT, 1)


def compter_oeuvres(mvt: dict) -> int:
    """Nombre d'œuvres transportées dans le mouvement (collection + hors collection).

    Un « mouvement » Navigart est un lot : il embarque plusieurs œuvres (champ
    `dos_nb_oeu_coll`, jusqu'à quelques centaines). On somme donc les masses œuvre
    par œuvre. Faute de domaine par œuvre, toutes sont estimées au forfait du domaine
    dominant. Défaut prudent : au moins une œuvre.
    """
    coll = int(mvt.get("dos_nb_oeu_coll") or 0)
    hcoll = int(mvt.get("dos_nb_oeu_hcoll") or 0)
    return max(1, coll + hcoll)


def choisir_mode(pays: str) -> str:
    """Impute le mode : routier en Europe, aérien hors d'Europe (défaut : routier)."""
    if not pays:
        return "Routier"
    return "Routier" if pays in PAYS_EUROPE else "Aérien"


def distance_ar_km(lat, lon, pays: str, mode: str):
    """Distance aller-retour siège <-> lieu, corrigée selon le mode.

    - Aérien : (orthodromie + 95 km) par vol, ×2 pour l'aller-retour (EN 16258).
    - Routier : orthodromie × facteur de détour national (Ballou et al. 2002), ×2.
    """
    if lat is None or lon is None:
        return None, None
    ortho = haversine_km(SIEGE_LAT, SIEGE_LON, float(lat), float(lon))
    if ortho < 0.5:                      # mouvement au siège même : pas de transport
        return 0.0, None
    if mode == "Aérien":
        return round(2 * (ortho + SUPPLEMENT_AERIEN_KM), 1), "EN 16258 (+95 km/vol)"
    facteur = FACTEUR_DETOUR.get(pays, FACTEUR_DETOUR_EUROPE)
    return round(2 * ortho * facteur, 1), f"détour ×{facteur} (Ballou 2002)"


def calculer_ligne(mvt: dict) -> dict:
    """Enrichit un mouvement avec distance corrigée, masse, mode, FE et émissions."""
    domaines = mvt.get("concerned_domains", "") or ""
    pays = mvt.get("location_country", "") or ""
    mode = choisir_mode(pays)
    distance_km, convention = distance_ar_km(
        mvt.get("latitude"), mvt.get("longitude"), pays, mode
    )
    nb_oeuvres = compter_oeuvres(mvt)
    masse_unitaire_kg = imputer_masse_kg(domaines)
    # Masse transportée = somme des œuvres du lot (masse unitaire × nombre d'œuvres).
    masse_totale_kg = round(masse_unitaire_kg * nb_oeuvres, 1)
    fe = FE_KG_PAR_TKM[mode]

    # Émissions ISO 14083 = distance (km) × masse (t) × FE (kgCO2e/t.km)
    emissions = None if distance_km is None else round((masse_totale_kg / 1000) * distance_km * fe, 3)

    return {
        "titre": mvt.get("title", ""),
        "domaine": domaines,
        "ville": mvt.get("location_city", ""),
        "pays": pays,
        "dates": mvt.get("exhibition_dates", ""),
        "distance_km_AR": distance_km,
        "convention_distance": convention or "",
        "nb_oeuvres": nb_oeuvres,
        "masse_kg_unitaire": masse_unitaire_kg,
        "masse_kg_totale": masse_totale_kg,
        "mode": mode,
        "fe_kg_tkm": fe,
        "emissions_kgCO2e": emissions,
    }


def ecrire_excel(lignes: list[dict]) -> float:
    """Écrit le classeur Excel : mouvements, cartographie ISO 14083, hypothèses sourcées."""
    classeur = Workbook()

    # --- Feuille 1 : mouvements enrichis --------------------------------------
    f1 = classeur.active
    f1.title = "Mouvements"
    colonnes = [
        "titre", "domaine", "ville", "pays", "dates", "distance_km_AR",
        "convention_distance", "nb_oeuvres", "masse_kg_unitaire", "masse_kg_totale",
        "mode", "fe_kg_tkm", "emissions_kgCO2e",
    ]
    f1.append(colonnes)
    for ligne in lignes:
        f1.append([ligne[c] for c in colonnes])

    total = sum(l["emissions_kgCO2e"] for l in lignes if l["emissions_kgCO2e"] is not None)
    f1.append([])
    f1.append(["TOTAL (kgCO2e)"] + [""] * (len(colonnes) - 2) + [round(total, 2)])

    # --- Feuille 2 : cartographie des « trous dans la raquette » (ISO 14083) ---
    f2 = classeur.create_sheet("Cartographie ISO 14083")
    f2.append(["Variable ISO 14083", "Statut dans Navigart", "Traitement dans la POC"])
    for ligne in [
        ("Coordonnées départ/arrivée", "Disponible (géocodé)", "Orthodromie (Haversine)"),
        ("Point de départ explicite", "Absent", "Hypothèse : retour au siège (Amiens)"),
        ("Distance réelle parcourue", "Absent", "Route : ortho × détour national (Ballou 2002) ; Air : +95 km/vol (EN 16258)"),
        ("Date précise de transport", "Partiel (dates d'expo)", "Non utilisée pour la distance"),
        ("Nombre d'œuvres du lot", "Disponible (dos_nb_oeu_coll/hcoll)", "Masse sommée sur toutes les œuvres du mouvement"),
        ("Masse de l'œuvre", "Absent", "Forfait par domaine dominant, appliqué à chaque œuvre du lot (hypothèse assumée)"),
        ("Masse du conditionnement", "Absent", "+30 % (convention GCC)"),
        ("Volume / poids volumétrique", "Absent", "Non traité (limite ; tarification au volume documentée par LP Art)"),
        ("Mode de transport", "Absent", "Imputé : Europe=routier / hors=aérien (Platform 2024 ; GCC 2022)"),
        ("Convoyage", "Absent", "Non mesurable"),
        ("Mutualisation / groupage", "Absent", "Non mesurable (le FE t.km alloue au prorata de la masse)"),
    ]:
        f2.append(list(ligne))

    # --- Feuille 3 : hypothèses et sources (traçabilité) -----------------------
    f3 = classeur.create_sheet("Hypotheses")
    f3.append(["Paramètre", "Valeur", "Source / statut"])
    f3.append(["Norme de calcul", "ISO 14083 (émissions = d × m × FE)", "ISO 14083:2023"])
    f3.append(["Principes de distance", "Distance réelle ; avion = orthodromie + 95 km/vol", "NF EN 16258 (2012)"])
    f3.append(["Détour routier France", "×1,65 (n=9, σ=0,46)", "Ballou et al. 2002 — facteur PAR PAYS (table en annexe)"])
    f3.append(["Détour routier Europe (défaut)", f"×{FACTEUR_DETOUR_EUROPE}", "Ballou et al. 2002 (n=199)"])
    f3.append(["Siège (lat, lon)", f"{SIEGE_LAT}, {SIEGE_LON}", "Confirmé (API Navigart)"])
    f3.append(["Trajectoire", "Aller-retour siège <-> lieu", "Hypothèse conservatrice, validée par le terrain"])
    for dom, m in MASSE_OEUVRE_NUE_KG.items():
        f3.append([f"Masse œuvre nue — {dom} (kg)", m, "Hypothèse de travail assumée (à affiner)"])
    f3.append(["Masse œuvre nue — défaut (kg)", MASSE_OEUVRE_DEFAUT_KG, "Hypothèse de travail assumée"])
    f3.append(["Masse du mouvement", "masse unitaire × nombre d'œuvres du lot", "Navigart : dos_nb_oeu_coll + dos_nb_oeu_hcoll (domaine dominant supposé pour tout le lot)"])
    f3.append(["Conditionnement", "+30 %", "Convention GCC Carbon Calculator (2024)"])
    f3.append(["Règle de mode", "Europe = routier / hors Europe = aérien", "Platform/Les Augures 2024 ; GCC 2022"])
    f3.append(["FE routier (kgCO2e/t.km)", FE_KG_PAR_TKM["Routier"], "ADEME Base Carbone v23.11 — Articulé 34-40 t, diesel 7 % bio (±70 %)"])
    f3.append(["FE aérien (kgCO2e/t.km)", FE_KG_PAR_TKM["Aérien"], "ADEME Base Carbone v23.11 — Avion cargo >100 t, >5000 km, 2023, AVEC traînées (±70 %)"])

    FICHIER_SORTIE.parent.mkdir(exist_ok=True)
    classeur.save(FICHIER_SORTIE)
    return total


def main() -> None:
    mouvements = charger_mouvements()
    lignes = [calculer_ligne(m) for m in mouvements]
    total = ecrire_excel(lignes)

    nb_aerien = sum(1 for l in lignes if l["mode"] == "Aérien")
    emis_aerien = sum(l["emissions_kgCO2e"] for l in lignes if l["mode"] == "Aérien" and l["emissions_kgCO2e"])
    dist_tot = sum(l["distance_km_AR"] for l in lignes if l["distance_km_AR"])
    nb_oeuvres = sum(l["nb_oeuvres"] for l in lignes)
    print(f"Mouvements traites        : {len(lignes)}")
    print(f"Oeuvres transportees      : {nb_oeuvres}")
    print(f"Dont aeriens (hors Europe): {nb_aerien}")
    print(f"Distance totale corrigee  : {round(dist_tot)} km")
    print(f"Total emissions estimees  : {round(total, 2)} kgCO2e (aerien AVEC trainees)")
    if total:
        print(f"  - part aerienne         : {round(100 * emis_aerien / total)} %")
    print(f"Classeur ecrit dans       : {FICHIER_SORTIE}")


if __name__ == "__main__":
    main()
